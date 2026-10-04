from __future__ import annotations

import re
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
FIG = ROOT / "figures"
OUT.mkdir(exist_ok=True)
FIG.mkdir(exist_ok=True)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_runs_with_bold(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def create_figures() -> None:
    metrics = ["Overall pass", "Attack success", "Privacy leakage", "Unauthorised tool use", "Approval violation", "Benign completion"]
    strict = [100, 0, 0, 0, 0, 100]
    permissive = [16.67, 100, 50, 33.33, 16.67, 100]
    x = range(len(metrics))
    fig, ax = plt.subplots(figsize=(10.5, 4.8))
    ax.bar([i - 0.19 for i in x], strict, width=0.38, label="Strict reference policy", color="#164E63")
    ax.bar([i + 0.19 for i in x], permissive, width=0.38, label="Permissive control", color="#D97706")
    ax.set_ylim(0, 110)
    ax.set_ylabel("Rate (%)")
    ax.set_xticks(list(x), metrics, rotation=24, ha="right")
    ax.grid(axis="y", alpha=0.22)
    ax.legend(frameon=False, ncol=2, loc="upper center")
    fig.tight_layout()
    fig.savefig(FIG / "baseline_comparison.png", dpi=220, bbox_inches="tight")
    plt.close(fig)

    labels = ["Risk categories", "Sectors", "Languages", "Scenarios"]
    values = [6, 5, 4, 120]
    fig, ax = plt.subplots(figsize=(8.4, 3.8))
    bars = ax.barh(labels, values, color=["#1D4ED8", "#2563EB", "#3B82F6", "#0F766E"])
    ax.set_xlabel("Count")
    ax.grid(axis="x", alpha=0.2)
    for bar, value in zip(bars, values):
        ax.text(value + (1.5 if value > 10 else 0.15), bar.get_y() + bar.get_height()/2, str(value), va="center", fontsize=10)
    ax.set_xlim(0, 132)
    fig.tight_layout()
    fig.savefig(FIG / "benchmark_coverage.png", dpi=220, bbox_inches="tight")
    plt.close(fig)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, before, after in [
        ("Title", 20, 0, 12),
        ("Heading 1", 14, 12, 6),
        ("Heading 2", 11.5, 9, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        p_pr = style.element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)

    footer = section.footer.paragraphs[0]
    footer.add_run("AgentGuardBench Technical Report  |  ")
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(90, 90, 90)
    add_page_number(footer)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("AgentGuardBench A Multilingual Benchmark for Privacy Security and Responsible Behaviour in AI Agents")
    p.paragraph_format.space_before = Pt(44)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Design and Deterministic Baseline Evaluation")
    run.bold = True
    run.font.size = Pt(14)

    for line in [
        "Joseph Arayemi",
        "GIIT Africa, Lagos, Nigeria",
        "ORCID 0009-0007-0776-7238",
        "Correspondence yemi@giitafrica.com",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Report DOI 10.5281/zenodo.23132194").bold = True
    p.add_run("\nSoftware DOI 10.5281/zenodo.23127436").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("3 October 2026")

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("Open technical report and reproducible research artifact")
    run.italic = True
    run.font.size = Pt(10)
    doc.add_page_break()


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    hdr = table.rows[0].cells
    for i, text in enumerate(rows[0]):
        hdr[i].text = text
        set_cell_shading(hdr[i], "1F4E78")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in hdr[i].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
        set_cell_border(hdr[i])
    for ridx, row in enumerate(rows[1:]):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = text
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ridx % 2:
                set_cell_shading(cells[i], "EEF4F8")
            set_cell_border(cells[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.size = Pt(9)
            if i > 0:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()


def build() -> Path:
    create_figures()
    lines = (ROOT / "manuscript.md").read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    add_cover(doc)

    i = 9  # Skip Markdown title and author block already represented on cover.
    in_code = False
    code_lines: list[str] = []
    figure_one_added = False
    figure_two_added = False
    current_section = ""
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.35)
                p.paragraph_format.right_indent = Inches(0.35)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Aptos Mono"
                run.font.size = Pt(9)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if line.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[-: |]+\|$", lines[i]):
                    table_lines.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            add_table(doc, table_lines)
            if not figure_one_added:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(FIG / "baseline_comparison.png"), width=Inches(6.55))
                cap = doc.add_paragraph("Figure 1. Deterministic baseline comparison across six evaluation metrics.")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].italic = True
                cap.runs[0].font.size = Pt(9)
                figure_one_added = True
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("## "):
            current_section = line[3:]
            heading = doc.add_paragraph(current_section, style="Heading 1")
            if line[3:] == "3 Benchmark Design" and not figure_two_added:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(FIG / "benchmark_coverage.png"), width=Inches(5.9))
                cap = doc.add_paragraph("Figure 2. Coverage dimensions in AgentGuardBench v0.1.1.")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.runs[0].italic = True
                cap.runs[0].font.size = Pt(9)
                figure_two_added = True
            i += 1
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
            i += 1
            continue
        if re.match(r"^\d+\. ", line):
            if current_section == "References":
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.22)
                p.paragraph_format.first_line_indent = Inches(-0.22)
                add_runs_with_bold(p, line)
            else:
                p = doc.add_paragraph(style="List Number")
                add_runs_with_bold(p, re.sub(r"^\d+\. ", "", line))
            i += 1
            continue
        if line.startswith("**Keywords:**"):
            p = doc.add_paragraph()
            add_runs_with_bold(p, line)
            i += 1
            continue
        if line.strip():
            p = doc.add_paragraph()
            add_runs_with_bold(p, line)
            if not line.startswith(("Source code", "Archived", "Concept", "License")):
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i += 1

    for section in doc.sections:
        section.footer_distance = Inches(0.32)
    doc.core_properties.title = "AgentGuardBench A Multilingual Benchmark for Privacy Security and Responsible Behaviour in AI Agents"
    doc.core_properties.author = "Joseph Arayemi"
    doc.core_properties.subject = "Responsible AI and AI agent security benchmark"
    doc.core_properties.keywords = "responsible AI, AI agents, security, privacy, prompt injection, benchmark"
    destination = OUT / "Joseph_Arayemi_AgentGuardBench_Technical_Report.docx"
    doc.save(destination)
    return destination


if __name__ == "__main__":
    print(build())
